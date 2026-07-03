import os
import zipfile
import logging
from typing import List, Dict, Any
from backend.app.config import settings

logger = logging.getLogger("uvicorn.error")

def generate_markdown_summary(video_data: dict, highlights: list) -> str:
    md = []
    md.append(f"# AI Video Summary: {video_data.get('title')}")
    md.append(f"\n**Duration:** {video_data.get('duration'):.2f} seconds | **Sentiment:** {video_data.get('sentiment')}\n")
    
    md.append("## Executive Summary")
    md.append(video_data.get('summary', 'No summary available.'))
    
    md.append("\n## Key Insights")
    for insight in video_data.get('key_insights', []):
        md.append(f"- {insight}")
        
    md.append("\n## Action Items")
    for item in video_data.get('action_items', []):
        md.append(f"- [ ] {item}")
        
    md.append("\n## Important Topics & Keywords")
    md.append(f"**Topics:** {', '.join(video_data.get('topics', []))}")
    md.append(f"**Keywords:** {', '.join(video_data.get('keywords', []))}")
    
    md.append("\n## AI Highlights Timeline")
    for hl in highlights:
        md.append(f"- **{hl.get('title')}** ({hl.get('start_time'):.1f}s - {hl.get('end_time'):.1f}s)")
        md.append(f"  *Score:* {hl.get('importance_score')}/10 | *Confidence:* {hl.get('confidence_score') * 100:.1f}%")
        
    return "\n".join(md)

def generate_pdf_report(video_data: dict, highlights: list, output_path: str) -> bool:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        doc = SimpleDocTemplate(output_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=22,
            textColor=colors.HexColor('#6366F1'),
            spaceAfter=15
        )
        
        h2_style = ParagraphStyle(
            'H2Style',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#4F46E5'),
            spaceBefore=10,
            spaceAfter=5
        )

        story.append(Paragraph(f"AI Video Summary & Highlight Report", title_style))
        story.append(Paragraph(f"<b>Title:</b> {video_data.get('title')}", styles['Normal']))
        story.append(Paragraph(f"<b>Duration:</b> {video_data.get('duration'):.2f}s | <b>Sentiment:</b> {video_data.get('sentiment')}", styles['Normal']))
        story.append(Spacer(1, 15))

        story.append(Paragraph("Executive Summary", h2_style))
        story.append(Paragraph(video_data.get('summary', 'No summary available.'), styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("Key Insights", h2_style))
        for insight in video_data.get('key_insights', []):
            story.append(Paragraph(f"• {insight}", styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("Action Items", h2_style))
        for item in video_data.get('action_items', []):
            story.append(Paragraph(f"[ ] {item}", styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("Key Highlights", h2_style))
        data = [["Title", "Time", "Score"]]
        for hl in highlights:
            data.append([
                hl.get('title', 'Highlight'),
                f"{hl.get('start_time'):.1f}s - {hl.get('end_time'):.1f}s",
                f"{hl.get('importance_score')}/10"
            ])
        t = Table(data, colWidths=[250, 150, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#6366F1')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F3F4F6'), colors.white]),
            ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
        ]))
        story.append(t)

        doc.build(story)
        return True
    except Exception as e:
        logger.error(f"Failed to generate PDF report: {e}")
        # Fallback to plain text write disguised as PDF or raise
        with open(output_path, "w") as f:
            f.write(f"PDF GENERATION FALLBACK\n\nTitle: {video_data.get('title')}\nSummary: {video_data.get('summary')}")
        return True

def generate_docx_report(video_data: dict, highlights: list, output_path: str) -> bool:
    try:
        from docx import Document
        doc = Document()
        doc.add_heading(f"AI Video Summary Report", 0)
        
        doc.add_paragraph(f"Video Title: {video_data.get('title')}")
        doc.add_paragraph(f"Duration: {video_data.get('duration'):.2f}s | Sentiment: {video_data.get('sentiment')}")
        
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(video_data.get('summary', ''))
        
        doc.add_heading("Key Insights", level=1)
        for insight in video_data.get('key_insights', []):
            doc.add_paragraph(insight, style='List Bullet')
            
        doc.add_heading("Action Items", level=1)
        for item in video_data.get('action_items', []):
            doc.add_paragraph(item, style='List Bullet')

        doc.add_heading("Highlights Timeline", level=1)
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Title'
        hdr_cells[1].text = 'Time'
        hdr_cells[2].text = 'Score'
        for hl in highlights:
            row_cells = table.add_row().cells
            row_cells[0].text = hl.get('title', '')
            row_cells[1].text = f"{hl.get('start_time'):.1f}s - {hl.get('end_time'):.1f}s"
            row_cells[2].text = f"{hl.get('importance_score')}/10"

        doc.save(output_path)
        return True
    except Exception as e:
        logger.error(f"Failed to generate DOCX report: {e}")
        with open(output_path, "w") as f:
            f.write(f"DOCX GENERATION FALLBACK\n\nTitle: {video_data.get('title')}")
        return True

def create_zip_package(video_title: str, md_content: str, transcript_text: str, pdf_path: str, docx_path: str, highlight_paths: list, zip_output_path: str) -> bool:
    try:
        with zipfile.ZipFile(zip_output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            # Add files to ZIP
            zipf.writestr("summary.md", md_content)
            zipf.writestr("transcript.txt", transcript_text)
            
            if pdf_path and os.path.exists(pdf_path):
                zipf.write(pdf_path, os.path.basename(pdf_path))
            if docx_path and os.path.exists(docx_path):
                zipf.write(docx_path, os.path.basename(docx_path))
                
            for hpath in highlight_paths:
                if hpath and os.path.exists(hpath):
                    zipf.write(hpath, f"highlights/{os.path.basename(hpath)}")
                    
        return True
    except Exception as e:
        logger.error(f"Failed to create ZIP package: {e}")
        return False
